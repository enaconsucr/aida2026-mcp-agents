"""Inter-annotator agreement between the authors' labels and the blind second annotator.
Uso: python3 acuerdo.py hoja_anotador.xlsx anotaciones_autores.json clave_mapeo.json
Reporta acuerdo crudo y kappa de Cohen por dimensión y para 'overall' (conjunción), con IC95% bootstrap
agrupado por pregunta (10 clusters) y detalle de discrepancias. Con n=20 los IC son amplios: es esperado.
Reportar ANTES de adjudicar; conservar ambas versiones."""
import json, sys, random
from openpyxl import load_workbook
DIMS = ["entity_grounding", "evidence_grounding", "causal_grounding"]
xl, auth, keyf = sys.argv[1:4]
key = json.load(open(keyf))
A = {(a["question_id"], a["model"]): a for a in json.load(open(auth))["annotations"]}
ws = load_workbook(xl, data_only=True).active
ann = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0]:
        assert all(v in (0, 1) for v in r[1:4]), f"Celda vacía o inválida en {r[0]}"
        ann[r[0]] = {"entity_grounding": r[1], "evidence_grounding": r[2], "causal_grounding": r[3], "c": r[4]}
rows = []
for cid, k in sorted(key.items()):
    a, b = A[(k["question_id"], k["model"])], ann[cid]
    ra = {d: a[d] for d in DIMS}; rb = {d: b[d] for d in DIMS}
    ra["overall"] = int(all(ra[d] for d in DIMS)); rb["overall"] = int(all(rb[d] for d in DIMS))
    rows.append((cid, k["question_id"], k["model"], ra, rb, b["c"]))
def kappa(x, y):
    n = len(x); po = sum(i == j for i, j in zip(x, y)) / n
    p1x, p1y = sum(x) / n, sum(y) / n; pe = p1x * p1y + (1 - p1x) * (1 - p1y)
    return po, (None if pe == 1 else (po - pe) / (1 - pe))
def boot(dim, B=5000):
    qs = sorted({r[1] for r in rows}); rng = random.Random(1); ks = []
    for _ in range(B):
        s = [r for q in (rng.choice(qs) for _ in qs) for r in rows if r[1] == q]
        k = kappa([r[3][dim] for r in s], [r[4][dim] for r in s])[1]
        if k is not None: ks.append(k)
    ks.sort(); return (ks[int(.025 * len(ks))], ks[int(.975 * len(ks))]) if len(ks) > 100 else (None, None)
print(f"{'dimensión':18} acuerdo  kappa   IC95% kappa (bootstrap por pregunta)")
for d in DIMS + ["overall"]:
    po, k = kappa([r[3][d] for r in rows], [r[4][d] for r in rows]); lo, hi = boot(d)
    f = lambda v: "n/d" if v is None else f"{v:.2f}"
    print(f"{d:18} {po:.2f}     {f(k):>5}  [{f(lo)}, {f(hi)}]")
print("\nDiscrepancias (autores vs anotador 2):")
for cid, q, m, ra, rb, c in rows:
    dd = [d for d in DIMS + ["overall"] if ra[d] != rb[d]]
    if dd: print(f"  {cid} {q} {m}: {', '.join(f'{d} {ra[d]}→{rb[d]}' for d in dd)}" + (f" | «{c}»" if c else ""))
